"""Tests for sofer._converters — normalization, sanitization, and dispatch.

Covers ``normalize_parquet_remote``, ``sanitize_sheet_name``,
``_write_parquet_table`` options, and ``convert_file_to_parquet`` for
csv/tsv/xlsx (single/multiple sheets), jsonl, parquet passthrough,
and the ``convert_to_parquet=false`` opt-out (caller-gated).
"""

from __future__ import annotations

from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from sofer import config
from sofer._converters import (
    CONVERTIBLE_SUFFIXES,
    _check_conversion_parity,
    _read_csv_raw_values,
    _write_parquet_table,
    convert_file_to_parquet,
    is_convertible_remote,
    leading_zero_columns,
    normalize_parquet_remote,
    sanitize_sheet_name,
)
from sofer._mirror import parquet_remote_for
from sofer.model import DatasetConfig, FileEntry

# ══════════════════════════════════════════════════════════════════════════════
#  normalize_parquet_remote
# ══════════════════════════════════════════════════════════════════════════════


class TestIsConvertibleRemote:
    """Issue #372: the single owner of the convertible-entry rule."""

    @pytest.mark.parametrize(
        ("remote", "expected"),
        [
            ("a.csv", True),
            ("a.CSV", True),
            ("a.tsv", True),
            ("a.xlsx", True),
            ("a.jsonl", True),
            ("a.JSoNl", True),
            ("folder\\a.CSV", True),
            ("nested/a.JSoNl", True),
            ("a.parquet", False),
            ("a.csv.gz", False),
            ("README", False),
        ],
    )
    def test_the_suffix_decides_after_folding_separators_and_case(
        self, remote: str, expected: bool
    ) -> None:
        assert is_convertible_remote(remote) is expected

    def test_a_recursive_entry_is_never_convertible(self) -> None:
        assert is_convertible_remote("a.csv", recursive=True) is False

    def test_conversion_and_schema_flags_are_not_part_of_the_rule(self) -> None:
        """Neither flag may filter the rule (issue #372 acceptance criterion).

        Case-fold collision validation considers every convertible remote even
        when ``convert_to_parquet`` or ``include_in_schema`` is false, so the
        predicate must take neither. This exercises every flag combination
        through ``FileEntry`` to prove the decision is unchanged.
        """
        for convert in (False, True):
            for include in (False, True):
                entry = FileEntry(
                    local=Path("unused"),
                    remote="a.csv",
                    convert_to_parquet=convert,
                    include_in_schema=include,
                )
                assert is_convertible_remote(entry.remote, recursive=entry.recursive) is True, (
                    convert,
                    include,
                )


class TestNormalizeParquetRemote:
    """PC-U02 normalization: lowercase, NFKD→ASCII, spaces→_, [^a-z0-9_./-]→_, collapse __+."""

    @pytest.mark.parametrize(
        ("inp", "expected"),
        [
            ("DATA GÖT Año.XLSX", "data_got_ano.xlsx"),
            ("My File 2024!.xlsx", "my_file_2024_.xlsx"),
            ("data/PROV/train.csv", "data/prov/train.csv"),
            ("data\\a\\train.csv", "data/a/train.csv"),
            ("a//b___c.csv", "a//b_c.csv"),
            ("A B", "a_b"),
            ("A-B", "a-b"),
            ("file.name-2024_v2.csv", "file.name-2024_v2.csv"),
            ("  leading  spaces  .csv", "_leading_spaces_.csv"),
        ],
    )
    def test_cases(self, inp: str, expected: str) -> None:
        assert normalize_parquet_remote(inp) == expected

    def test_parquet_key_derivation_via_mirror(self) -> None:
        """Combined ``parquet_remote_for`` + ``normalize`` yields the staged key."""
        assert (
            normalize_parquet_remote(parquet_remote_for("DATA GÖT Año.XLSX"))
            == "data_got_ano.parquet"
        )
        assert (
            normalize_parquet_remote(parquet_remote_for("My File 2024!.xlsx"))
            == "my_file_2024_.parquet"
        )

    def test_preserves_dots_and_slashes(self) -> None:
        assert normalize_parquet_remote("a/b.c/d_e-f.csv") == "a/b.c/d_e-f.csv"

    def test_collapses_underscores(self) -> None:
        assert normalize_parquet_remote("a__b___c.csv") == "a_b_c.csv"
        # spaces become underscores then collapse
        assert normalize_parquet_remote("a  b.csv") == "a_b.csv"

    def test_accent_stripping(self) -> None:
        assert normalize_parquet_remote("café.csv") == "cafe.csv"
        assert normalize_parquet_remote("naïve.csv") == "naive.csv"

    def test_backslash_to_slash(self) -> None:
        assert normalize_parquet_remote("data\\PROV\\train.csv") == "data/prov/train.csv"

    def test_a_space_vs_dash_distinct(self) -> None:
        """``A B`` → ``a_b`` and ``A-B`` → ``a-b`` must not collide."""
        assert normalize_parquet_remote("A B") != normalize_parquet_remote("A-B")
        assert normalize_parquet_remote("A B") == "a_b"
        assert normalize_parquet_remote("A-B") == "a-b"


# ══════════════════════════════════════════════════════════════════════════════
#  sanitize_sheet_name
# ══════════════════════════════════════════════════════════════════════════════


class TestSanitizeSheetName:
    """Sheet sanitization: lowercase, NFKD→ASCII, spaces→_, narrow alphabet, collapse, fallback."""

    @pytest.mark.parametrize(
        ("inp", "expected"),
        [
            ("Ventas 2024!", "ventas_2024"),
            ("Sheet", "sheet"),
            ("  __Hello--World__  ", "hello--world"),
            ("GÖT", "got"),
            ("", "sheet"),
            ("!!!", "sheet"),
            ("A  B", "a_b"),
            ("A-B", "a-b"),
        ],
    )
    def test_cases(self, inp: str, expected: str) -> None:
        assert sanitize_sheet_name(inp) == expected

    def test_dedup_via_converter(self, tmp_path: Path) -> None:
        """Two sheets sanitizing to the same key get ``_2`` suffix via the converter."""
        import openpyxl

        xlsx = tmp_path / "dup.xlsx"
        wb = openpyxl.Workbook()
        ws1 = wb.active
        assert ws1 is not None
        ws1.title = "A B"
        ws1["A1"] = "v"
        ws1["A2"] = "1"
        ws2 = wb.create_sheet(title="A_B")
        ws2["A1"] = "v"
        ws2["A2"] = "2"
        wb.save(xlsx)

        staging = tmp_path / "staging"
        result = convert_file_to_parquet(xlsx, staging)
        # Keys are the sanitized stems with dedup: a_b and a_b_2
        assert "dup__a_b" in result
        assert "dup__a_b_2" in result
        assert len(result) == 2


# ══════════════════════════════════════════════════════════════════════════════
#  _write_parquet_table
# ══════════════════════════════════════════════════════════════════════════════


class TestWriteParquetTable:
    """Writer: compression, row group size, null→string cast."""

    def test_null_column_cast_to_string(self, tmp_path: Path) -> None:
        """All-null columns are cast to string so they render correctly."""
        table = pa.table({"a": [None, None], "b": [1, 2]})
        out = _write_parquet_table(table, tmp_path, "test_null")
        assert out.is_file()
        read = pq.read_table(out)
        assert pa.types.is_string(read.schema.field("a").type)

    def test_compression_and_row_group(self, tmp_path: Path) -> None:
        """Written file uses configured compression and row group size."""
        table = pa.table({"x": list(range(10))})
        out = _write_parquet_table(table, tmp_path, "test_comp")
        assert out.is_file()
        pf = pq.ParquetFile(out)
        # Verify compression via metadata
        row_group = pf.metadata.row_group(0)
        col_chunk = row_group.column(0)
        # Compression should be ZSTD as per config default
        assert col_chunk.compression.lower() == config.PARQUET_COMPRESSION.lower()

    def test_creates_parent_dirs(self, tmp_path: Path) -> None:
        table = pa.table({"x": [1]})
        nested = tmp_path / "a" / "b"
        out = _write_parquet_table(table, nested, "nested")
        assert out.is_file()
        assert out.parent == nested


# ══════════════════════════════════════════════════════════════════════════════
#  convert_file_to_parquet dispatcher
# ══════════════════════════════════════════════════════════════════════════════


class TestConvertFileToParquet:
    """Dispatcher: csv/tsv/xlsx (single/multi), jsonl, parquet passthrough."""

    def test_csv(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("a;b\n1;2\n3;4\n", encoding="utf-8-sig")
        staging = tmp_path / "staging"
        result = convert_file_to_parquet(csv, staging)
        assert "data" in result
        assert result["data"].is_file()
        assert pq.read_table(result["data"]).num_rows == 2

    def test_tsv(self, tmp_path: Path) -> None:
        tsv = tmp_path / "data.tsv"
        tsv.write_text("a\tb\n1\t2\n", encoding="utf-8-sig")
        staging = tmp_path / "staging"
        result = convert_file_to_parquet(tsv, staging)
        assert "data" in result
        assert pq.read_table(result["data"]).num_rows == 1

    def test_xlsx_single_sheet(self, tmp_path: Path) -> None:
        import openpyxl

        xlsx = tmp_path / "single.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Data"
        ws["A1"] = "a"
        ws["B1"] = "b"
        ws["A2"] = 1
        ws["B2"] = 2
        wb.save(xlsx)

        staging = tmp_path / "staging"
        result = convert_file_to_parquet(xlsx, staging)
        assert "single" in result
        assert pq.read_table(result["single"]).num_rows == 1

    def test_xlsx_multi_sheet(self, tmp_path: Path) -> None:
        import openpyxl

        xlsx = tmp_path / "multi.xlsx"
        wb = openpyxl.Workbook()
        ws1 = wb.active
        assert ws1 is not None
        ws1.title = "Ventas 2024!"
        ws1["A1"] = "x"
        ws1["A2"] = "1"
        ws2 = wb.create_sheet(title="Other")
        ws2["A1"] = "y"
        ws2["A2"] = "2"
        wb.save(xlsx)

        staging = tmp_path / "staging"
        result = convert_file_to_parquet(xlsx, staging)
        assert len(result) == 2
        assert "multi__ventas_2024" in result
        assert "multi__other" in result
        for p in result.values():
            assert p.is_file()

    def test_jsonl(self, tmp_path: Path) -> None:
        j = tmp_path / "data.jsonl"
        j.write_text('{"a": 1, "b": "x"}\n{"a": 2, "b": "y"}\n', encoding="utf-8")
        staging = tmp_path / "staging"
        result = convert_file_to_parquet(j, staging)
        assert "data" in result
        assert pq.read_table(result["data"]).num_rows == 2

    def test_parquet_passthrough(self, tmp_path: Path) -> None:
        p = tmp_path / "orig.parquet"
        pq.write_table(pa.table({"x": [1, 2]}), p)
        staging = tmp_path / "staging"
        result = convert_file_to_parquet(p, staging)
        assert result == {}

    def test_unknown_suffix(self, tmp_path: Path) -> None:
        txt = tmp_path / "notes.txt"
        txt.write_text("hello", encoding="utf-8")
        staging = tmp_path / "staging"
        result = convert_file_to_parquet(txt, staging)
        assert result == {}

    def test_convert_to_parquet_false_opt_out_via_prepare(self, tmp_path: Path) -> None:
        """Caller-gated opt-out: prepare with convert_to_parquet=False stages CSV as-is."""
        from sofer.model import DatasetConfig, FileEntry
        from sofer.prepare import prepare

        csv = tmp_path / "data.csv"
        csv.write_text("a;b\n1;2\n", encoding="utf-8-sig")
        cfg = DatasetConfig(
            name="test",
            repo_id="user/test",
            files=[FileEntry(local=csv, remote="data.csv", convert_to_parquet=False)],
            _base_dir=tmp_path,
        )
        out = tmp_path / "build"
        rc = prepare(cfg, out)
        assert rc == 0
        assert (out / "data.csv").is_file()
        assert not (out / "data.parquet").exists()

    def test_convertible_suffixes_constant(self) -> None:
        assert ".csv" in CONVERTIBLE_SUFFIXES
        assert ".tsv" in CONVERTIBLE_SUFFIXES
        assert ".xlsx" in CONVERTIBLE_SUFFIXES
        assert ".jsonl" in CONVERTIBLE_SUFFIXES
        assert ".parquet" not in CONVERTIBLE_SUFFIXES


class TestJsonlFallback:
    """The JSONL fallback path is exercised and its real behaviour asserted."""

    def test_from_pylist_fallback_uses_first_record_schema(
        self, tmp_path: Path, monkeypatch, capsys
    ) -> None:
        """Forced fallback: the first record defines the schema; later-only keys drop.

        ``pyarrow.json.read_json`` is forced to raise so the ``json`` +
        ``pa.Table.from_pylist`` fallback runs. On the installed pyarrow,
        ``from_pylist`` infers the schema from the first record, so a key that
        appears only in a later record is not unioned in — the fallback's code
        comment is corrected to match this. The key-union parity check is
        advisory and the conversion still succeeds (issue #205 F2).
        """

        def _raise(*_args: object, **_kwargs: object) -> object:
            raise RuntimeError("forced fast-path failure")

        import pyarrow.json as pj

        monkeypatch.setattr(pj, "read_json", _raise)

        j = tmp_path / "data.jsonl"
        j.write_text('{"a": 1, "c": "first"}\n{"a": 2, "b": "later"}\n', encoding="utf-8")
        staging = tmp_path / "staging"
        result = convert_file_to_parquet(j, staging)
        assert "data" in result
        table = pq.read_table(result["data"])
        assert table.num_rows == 2
        assert table.column_names == ["a", "c"]
        assert "b" not in table.column_names
        assert "JSONL key mismatch" in capsys.readouterr().out


class TestCsvDialectResolution:
    """Declared dialect vs fallback resolution (PC-U01, E2)."""

    def _cfg_with_declared(self, delimiter: str) -> DatasetConfig:
        """Programmatic cfg carrying the declared-keys presence signal."""
        return DatasetConfig(
            name="t",
            repo_id="user/t",
            files=[],
            csv_delimiter=delimiter,
            declared_meta_keys=frozenset({"csv_delimiter"}),
        )

    def test_declared_semicolon_wins(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("a;b\n1;2\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging", self._cfg_with_declared(";"))
        assert result
        assert pq.read_table(result["data"]).column_names == ["a", "b"]

    def test_declared_comma_wins(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("a,b\n1,2\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging", self._cfg_with_declared(","))
        assert result
        assert pq.read_table(result["data"]).column_names == ["a", "b"]

    def test_declared_pipe_outside_sniff_set_wins(self, tmp_path: Path) -> None:
        """A declared ``|`` is the only way to keep the true columns.

        ``;`` and ``,`` are in ``config.SNIFF_DELIMITERS``, so the sniff picks
        the same delimiter and those cases pass whether or not the declaration
        is honoured. ``|`` is outside the sniff set, so a collapse to one column
        is the only alternative — the case cannot pass by coincidence
        (issue #205 F3).
        """
        csv = tmp_path / "data.csv"
        csv.write_text("col1|col2|col3\nv1|v2|v3\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging", self._cfg_with_declared("|"))
        assert result
        assert pq.read_table(result["data"]).column_names == ["col1", "col2", "col3"]

    def test_nothing_declared_falls_back_to_sniff(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("a;b\n1;2\n", encoding="utf-8-sig")
        cfg = DatasetConfig(name="t", repo_id="user/t", files=[])
        assert "csv_delimiter" not in cfg.declared_meta_keys
        result = convert_file_to_parquet(csv, tmp_path / "staging", cfg)
        assert result
        assert pq.read_table(result["data"]).column_names == ["a", "b"]

    def test_programmatic_value_without_presence_signal_falls_back_to_sniff(
        self, tmp_path: Path
    ) -> None:
        """A programmatic dialect with no ``declared_meta_keys`` does not override.

        The presence signal is the contract of record (issue #181 design D1;
        issue #205 F5): ``from_toml`` alone populates ``declared_meta_keys``, so
        a library consumer passing ``csv_delimiter=`` directly takes the sniff
        fallback. This test pins that accepted boundary — a ``|`` file collapses
        to one column because the declaration is not observed.
        """
        csv = tmp_path / "data.csv"
        csv.write_text("col1|col2|col3\nv1|v2|v3\n", encoding="utf-8-sig")
        cfg = DatasetConfig(name="t", repo_id="user/t", files=[], csv_delimiter="|")
        assert "csv_delimiter" not in cfg.declared_meta_keys
        result = convert_file_to_parquet(csv, tmp_path / "staging", cfg)
        assert result
        assert pq.read_table(result["data"]).num_columns == 1

    def test_declared_encoding_reaches_python_parity_read(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_bytes("name\nJos\u00e9\n".encode("cp1252"))
        cfg = DatasetConfig(
            name="t",
            repo_id="user/t",
            files=[],
            csv_encoding="cp1252",
            declared_meta_keys=frozenset({"csv_encoding"}),
        )
        result = convert_file_to_parquet(csv, tmp_path / "staging", cfg)
        assert result
        assert pq.read_table(result["data"]).column("name").to_pylist() == ["Jos\u00e9"]

    def test_read_csv_raw_values_honours_encoding(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_bytes("name\nJos\u00e9\n".encode("cp1252"))
        assert _read_csv_raw_values(csv, ";", "cp1252") == (["name"], [["Jos\u00e9"]])
        assert _read_csv_raw_values(csv, ";") is None

    def test_collapse_guard_fires_when_raw_header_carries_resolved_delimiter(
        self, tmp_path: Path, capsys
    ) -> None:
        """A one-column table whose raw header still carries the delimiter is a mis-split.

        The guard keys on the *resolved* delimiter (design D3) — the reading that
        keeps the D2 disagreement case silent — so a hand-built one-column table
        is the only way to reach the branch with the delimiter genuinely present.
        """
        csv = tmp_path / "data.csv"
        csv.write_text("a;b\n1\n", encoding="utf-8-sig")
        table = pa.table({"a;b": [1]})
        assert table.num_columns == 1
        ok, _warnings = _check_conversion_parity(csv, ";", table, config.CSV_ENCODING)
        assert ok is False
        assert "single column but the raw header contains" in capsys.readouterr().out

    def test_collapse_guard_silent_when_declared_delimiter_absent(self, tmp_path: Path) -> None:
        import pyarrow.csv as pc

        csv = tmp_path / "data.csv"
        csv.write_text("a,b\n1,2\n", encoding="utf-8-sig")
        table = pc.read_csv(csv, parse_options=pc.ParseOptions(delimiter=";"))
        assert table.num_columns == 1
        ok, _warnings = _check_conversion_parity(csv, ";", table, config.CSV_ENCODING)
        assert ok is True


class TestConvertFileToParquetRetargetedArms:
    """Arms retargeted from the removed prepare conversion cluster (D4)."""

    def test_csv_type_inference_and_values(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("id;name;age\n1;Alice;30\n2;Bob;25\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging")
        table = pq.read_table(result["data"])
        assert table.num_rows == 2
        assert table.column_names == ["id", "name", "age"]
        assert table.column("name").to_pylist() == ["Alice", "Bob"]
        assert table.schema.field("id").type == pa.int64()

    def test_corrupt_csv_returns_empty_and_warns(self, tmp_path: Path, capsys) -> None:
        csv = tmp_path / "corrupt.csv"
        csv.write_bytes(bytes([0, 1, 2, 255, 254]))
        result = convert_file_to_parquet(csv, tmp_path / "staging")
        assert result == {}
        assert "conversion failed" in capsys.readouterr().out.lower()

    def test_missing_csv_returns_empty(self, tmp_path: Path) -> None:
        assert convert_file_to_parquet(tmp_path / "ghost.csv", tmp_path / "staging") == {}

    def test_parity_failure_returns_empty(self, tmp_path: Path) -> None:
        csv = tmp_path / "parity.csv"
        csv.write_text("a;b\n1;2\n\n\n", encoding="utf-8")
        assert convert_file_to_parquet(csv, tmp_path / "staging") == {}

    def test_oversized_shard_warns(self, tmp_path: Path, monkeypatch, capsys) -> None:
        csv = tmp_path / "big.csv"
        csv.write_text("a;b\n" + "1;2\n" * 2000, encoding="utf-8")
        monkeypatch.setattr(config, "PARQUET_SHARD_WARNING_MB", 0.000001)
        result = convert_file_to_parquet(csv, tmp_path / "staging")
        assert result
        assert "Consider sharding into" in capsys.readouterr().out

    def test_value_parity_warns_per_column(self, tmp_path: Path, capsys) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("zip;sku\n01234;00099\n05678;00100\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging")
        assert result
        out = capsys.readouterr().out
        assert "zip" in out and "sku" in out

    def test_value_parity_silent_when_values_match(self, tmp_path: Path, capsys) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("name;age\nAlice;30\nBob;25\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging")
        assert result
        assert "altered" not in capsys.readouterr().out.lower()

    def test_all_null_column_via_conversion_cast_to_string(self, tmp_path: Path) -> None:
        csv = tmp_path / "data.csv"
        csv.write_text("name;empty_col\nAlice;\nBob;\n", encoding="utf-8-sig")
        result = convert_file_to_parquet(csv, tmp_path / "staging")
        table = pq.read_table(result["data"])
        col_type = table.schema.field("empty_col").type
        assert pa.types.is_string(col_type) or pa.types.is_large_string(col_type)


# ══════════════════════════════════════════════════════════════════════════════
#  Issue #337 — declared dtype control and zero-padded code protection
# ══════════════════════════════════════════════════════════════════════════════


class TestLeadingZeroProbe:
    """The probe that decides which columns a numeric conversion would alter."""

    def test_zero_padded_code_column_is_detected(self, tmp_path: Path) -> None:
        csv = tmp_path / "codes.csv"
        csv.write_text("ID;NAME\n001;a\n02;b\n5;c\n", encoding="utf-8")
        assert leading_zero_columns(csv, ";") == ["ID"]

    def test_plain_integer_column_is_not_detected(self, tmp_path: Path) -> None:
        csv = tmp_path / "plain.csv"
        csv.write_text("ID;NAME\n1;a\n2;b\n", encoding="utf-8")
        assert leading_zero_columns(csv, ";") == []

    def test_text_with_an_incidental_zero_token_is_not_detected(self, tmp_path: Path) -> None:
        """A column that is not all digits is never forced to string."""
        csv = tmp_path / "mixed.csv"
        csv.write_text("NOTE\n007 licence\nsee 001 above\n", encoding="utf-8")
        assert leading_zero_columns(csv, ";") == []

    def test_empty_cells_do_not_disqualify_a_code_column(self, tmp_path: Path) -> None:
        csv = tmp_path / "gaps.csv"
        csv.write_text("ID\n001\n\n02\n", encoding="utf-8")
        assert leading_zero_columns(csv, ";") == ["ID"]

    def test_probe_stops_at_the_configured_row_budget(self, tmp_path: Path) -> None:
        """A code that only appears past the budget is not probed (bounded cost)."""
        rows = ["ID"] + ["1"] * 5 + ["001"]
        csv = tmp_path / "late.csv"
        csv.write_text("\n".join(rows) + "\n", encoding="utf-8")
        assert leading_zero_columns(csv, ";", probe_rows=5) == []
        assert leading_zero_columns(csv, ";", probe_rows=10) == ["ID"]

    def test_unreadable_file_returns_empty(self, tmp_path: Path) -> None:
        assert leading_zero_columns(tmp_path, ";") == []

    def test_header_only_file_returns_empty(self, tmp_path: Path) -> None:
        csv = tmp_path / "empty.csv"
        csv.write_text("ID;NAME\n", encoding="utf-8")
        assert leading_zero_columns(csv, ";") == []


class TestColumnTypeControl:
    """#337: the declared dtype and the automatic code protection."""

    def test_declared_column_types_forces_string(self, tmp_path: Path) -> None:
        """A declared column reaches the Parquet as its source string form."""
        csv = tmp_path / "declared.csv"
        csv.write_text("ID;NAME\n1;a\n2;b\n", encoding="utf-8")
        staging = tmp_path / "stage"

        result = convert_file_to_parquet(csv, staging, None, column_types={"ID": "string"})

        table = pq.read_table(result["declared"])
        assert table.schema.field("ID").type == pa.string()
        assert table.column("ID").to_pylist() == ["1", "2"]

    def test_zero_padded_column_is_protected_without_config(self, tmp_path: Path, capsys) -> None:
        """T1: the source form survives without any declaration, and says so."""
        csv = tmp_path / "codes.csv"
        csv.write_text("ID;NAME\n001;a\n02;b\n", encoding="utf-8")
        staging = tmp_path / "stage"

        result = convert_file_to_parquet(csv, staging, None)

        table = pq.read_table(result["codes"])
        assert table.schema.field("ID").type == pa.string()
        assert table.column("ID").to_pylist() == ["001", "02"]
        out = capsys.readouterr().out
        assert "carries zero-padded codes" in out
        assert 'column_types = { ID = "string" }' in out

    def test_unknown_declared_type_is_refused_naming_the_set(self) -> None:
        """T2: the vocabulary is closed; a typo fails closed."""
        from sofer._converters import column_type_datatype

        with pytest.raises(ValueError, match="accepted: 'string'"):
            column_type_datatype("int64")

    def test_unsupported_declared_type_fails_closed(self, tmp_path: Path, capsys) -> None:
        """An unvalidated type never converts: the original CSV is staged."""
        csv = tmp_path / "declared.csv"
        csv.write_text("ID;NAME\n1;a\n", encoding="utf-8")

        result = convert_file_to_parquet(
            csv, tmp_path / "stage", None, column_types={"ID": "int64"}
        )

        assert result == {}
        out = capsys.readouterr().out
        assert "conversion failed" in out
        assert "accepted: 'string'" in out

    def test_residual_alteration_raises_instead_of_shipping(self, tmp_path: Path) -> None:
        """T3: an alteration no forced dtype covers is fatal (D4)."""
        from sofer._converters import ValueAlterationError

        csv = tmp_path / "plus.csv"
        csv.write_text("V;NAME\n+5;a\n+6;b\n", encoding="utf-8")

        with pytest.raises(ValueAlterationError) as excinfo:
            convert_file_to_parquet(csv, tmp_path / "stage", None)

        assert excinfo.value.filename == "plus.csv"
        assert [name for name, _ in excinfo.value.columns] == ["V"]

    def test_declared_type_wins_over_the_probe(self, tmp_path: Path, capsys) -> None:
        """A declared column is not re-announced as an automatic decision."""
        csv = tmp_path / "codes.csv"
        csv.write_text("ID;NAME\n001;a\n", encoding="utf-8")

        convert_file_to_parquet(csv, tmp_path / "stage", None, column_types={"ID": "string"})

        assert "carries zero-padded codes" not in capsys.readouterr().out

    def test_dataset_without_codes_keeps_the_pre_change_read(self, tmp_path: Path) -> None:
        """No code column and no declaration → the table is inferred as before."""
        csv = tmp_path / "plain.csv"
        csv.write_text("ID;NAME\n1;a\n", encoding="utf-8")

        result = convert_file_to_parquet(csv, tmp_path / "stage", None)

        table = pq.read_table(result["plain"])
        assert table.schema.field("ID").type == pa.int64()
